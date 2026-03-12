"""
ml/extractor.py
Extracts plain text from PDF and DOCX resume files.
"""
import io
import logging

logger = logging.getLogger(__name__)


def extract_text(file_obj) -> str:
    """
    Extract raw text from an uploaded file object.
    Supports PDF (.pdf) and Word (.docx) formats.
    Returns empty string on failure.
    """
    filename = getattr(file_obj, 'name', '').lower()

    try:
        if filename.endswith('.pdf'):
            return _extract_pdf(file_obj)
        elif filename.endswith('.docx'):
            return _extract_docx(file_obj)
        else:
            # Try to read as plain text
            content = file_obj.read()
            return content.decode('utf-8', errors='ignore')
    except Exception as e:
        logger.error(f"Text extraction failed for {filename}: {e}")
        return ''


def _extract_pdf(file_obj) -> str:
    """Extract text from PDF using pypdf."""
    from pypdf import PdfReader

    file_obj.seek(0)
    reader = PdfReader(io.BytesIO(file_obj.read()))
    pages_text = []
    for page in reader.pages:
        text = page.extract_text()
        if text:
            pages_text.append(text)
    return '\n'.join(pages_text)


def _extract_docx(file_obj) -> str:
    """Extract text from DOCX using python-docx."""
    from docx import Document

    file_obj.seek(0)
    doc = Document(io.BytesIO(file_obj.read()))
    paragraphs = [para.text for para in doc.paragraphs if para.text.strip()]
    # Also grab table cells
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    paragraphs.append(cell.text.strip())
    return '\n'.join(paragraphs)
